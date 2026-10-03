"""Run with: python3 -m unittest discover -s skills/stolen-bell-tipout/tests"""
import json
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "scripts"))

try:
    from openpyxl import load_workbook
    from fill_cashout import fill
except ImportError:  # pragma: no cover
    fill = None

try:
    import formulas  # evaluates the form's real Excel formulas
except ImportError:  # pragma: no cover
    formulas = None

from tipout import compute_day  # noqa: E402

EXAMPLE = json.loads((HERE.parent / "examples" / "dinner_example.json").read_text())


def one_period(staff, card=1000.0, food=2000.0, net=4000.0, cash=0):
    return {"business_date": "2026-09-27", "periods": {"Dinner": {
        "pool_card_tips": card, "auto_gratuity": 0, "cash_tips_manual": cash,
        "gross_food_sales": food, "net_sales": net, "staff": staff}}}


@unittest.skipIf(fill is None, "openpyxl not installed")
class FillCashOut(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.out = Path(self.tmp.name) / "cashout.xlsx"

    def tearDown(self):
        self.tmp.cleanup()

    def test_fills_one_sheet_per_period(self):
        summary = fill(EXAMPLE, self.out, "S. Thornton")
        self.assertEqual(summary["problems"], [])
        wb = load_workbook(self.out)
        self.assertEqual(wb.sheetnames, ["Lunch", "Dinner"])
        ws = wb["Dinner"]
        self.assertEqual((ws["B4"].value, ws["E4"].value, ws["B5"].value), ("2026-09-27", "Dinner", "S. Thornton"))
        self.assertEqual((ws["B9"].value, ws["B10"].value, ws["B14"].value), (9500, 4200, 1800))
        self.assertEqual([(ws[f"B{r}"].value, ws[f"E{r}"].value) for r in (23, 24, 25)],
                         [("Priya (Host)", "Host"), ("Tomas (Barback)", "Barback"), (None, None)])
        self.assertEqual([ws[f"B{r}"].value for r in (30, 31, 32, 33)], ["Mara (Bartender)", "Devon (Server)", "Sam (Server)", None])
        self.assertEqual(ws["B26"].value, "=B16-B19-B21")  # formulas are left intact

    def evaluate(self):
        solution = formulas.ExcelModel().loads(str(self.out)).finish().calculate()

        def cell(sheet, ref):
            key = next(k for k in solution if k.upper().endswith(f"]{sheet.upper()}'!{ref}"))
            return Decimal(str(solution[key].value[0][0]))
        return cell

    @unittest.skipIf(formulas is None, "formulas not installed")
    def test_the_forms_own_formulas_agree_with_the_engine(self):
        fill(EXAMPLE, self.out, None)
        cell = self.evaluate()

        day = compute_day(EXAMPLE)
        for r in day["periods"]:
            self.assertAlmostEqual(cell(r.period, "B19"), r.kitchen, places=2)
            self.assertAlmostEqual(cell(r.period, "B21"), r.support_total, places=2)
            self.assertEqual(cell(r.period, "B45"), 0)
        self.assertLess(abs(cell("Dinner", "E30") - day["periods"][1].tipped_payouts["Mara (Bartender)"]), Decimal("0.01"))
        self.assertLess(abs(cell("Dinner", "D30") - day["periods"][1].hourly_rate), Decimal("0.01"))

    @unittest.skipIf(formulas is None, "formulas not installed")
    def test_two_hosts_get_the_one_role_rate_on_the_form(self):
        doc = one_period([{"name": "H1", "role": "Host", "hours": 4}, {"name": "H2", "role": "Host", "hours": 4},
                          {"name": "S", "role": "Server", "hours": 6}])
        summary = fill(doc, self.out, None)
        self.assertEqual(summary["problems"], [])
        cell = self.evaluate()
        self.assertEqual(cell("Dinner", "B20"), Decimal("0.015"))
        self.assertAlmostEqual(cell("Dinner", "B21"), Decimal("60"), places=2)
        self.assertEqual(cell("Dinner", "B45"), 0)

    @unittest.skipIf(formulas is None, "formulas not installed")
    def test_host_and_barback_with_a_third_support_person(self):
        doc = one_period([{"name": "H1", "role": "Host", "hours": 4}, {"name": "H2", "role": "Host", "hours": 2},
                          {"name": "B", "role": "Barback", "hours": 6}, {"name": "S", "role": "Server", "hours": 6}])
        summary = fill(doc, self.out, None)
        self.assertEqual(summary["problems"], [])
        cell = self.evaluate()
        self.assertEqual(cell("Dinner", "B20"), Decimal("0.0225"))
        self.assertAlmostEqual(cell("Dinner", "D25"), Decimal("45"), places=2)  # 90 support x 6/12 h

    def test_refuses_when_support_staff_do_not_fit(self):
        staff = [{"name": f"H{i}", "role": "Host", "hours": 2} for i in range(4)]
        summary = fill(one_period(staff + [{"name": "S", "role": "Server", "hours": 6}]), self.out, None)
        self.assertIsNone(summary["written"])
        self.assertTrue(any("only has 3 rows" in p for p in summary["problems"]))
        self.assertFalse(self.out.exists())

    def test_refuses_on_due_back(self):
        doc = one_period([{"name": "S", "role": "Server", "hours": 6}], card=100, food=2000)
        summary = fill(doc, self.out, None)
        self.assertTrue(any("DUE BACK" in p for p in summary["problems"]))

    def test_refuses_when_staff_do_not_fit(self):
        staff = [{"name": f"S{i}", "role": "Server", "hours": 4} for i in range(11)]
        summary = fill(one_period(staff), self.out, None)
        self.assertTrue(any("only has 10 rows" in p for p in summary["problems"]))

    def test_one_row_per_person_when_listed_twice(self):
        doc = one_period([{"name": "Sam", "role": "Bartender", "hours": 3}, {"name": "Sam", "role": "Server", "hours": 3},
                          {"name": "Ana", "role": "Server", "hours": 6}], food=0, net=0)
        summary = fill(doc, self.out, None)
        self.assertEqual(summary["problems"], [])
        ws = load_workbook(self.out)["Dinner"]
        self.assertEqual([(ws[f"B{r}"].value, ws[f"C{r}"].value) for r in (30, 31)], [("Sam", 6), ("Ana", 6)])


if __name__ == "__main__":
    unittest.main()
