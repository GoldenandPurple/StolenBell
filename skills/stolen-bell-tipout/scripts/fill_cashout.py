#!/usr/bin/env python3
"""
Fill Steph's Cash Out form (forms/CashOut_Form.xlsx) from a tip-out input document.

Writes one sheet per service period with the yellow input cells filled in; the
form's own formulas then calculate the payouts as they always have. Before
saving, it checks that those formulas will land on the same figures as the
tip-out engine (scripts/tipout.py). If they won't, nothing is written and the
differences are reported, because a form that disagrees with the rules must not
be the thing someone pays from.

Usage:
  python3 scripts/fill_cashout.py <input.json> --out <CashOut_YYYY-MM-DD.xlsx> [--completed-by NAME]

Prints a JSON summary. Exit code 0 = written, 2 = not written (see "problems").
Needs openpyxl.
"""

from __future__ import annotations

import argparse
import json
import sys
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from tipout import SUPPORT_ROLES, TIPPED_ROLES, compute_day  # noqa: E402

FORM = Path(__file__).resolve().parent.parent / "forms" / "CashOut_Form.xlsx"
TEMPLATE_SHEET = "Cash Out"
SUPPORT_ROWS = [23, 24, 25]
TIPPED_ROWS = list(range(30, 40))
TOLERANCE = Decimal("0.01")  # the form doesn't round to the cent; the engine does


def _d(value) -> Decimal:
    return Decimal(str(value or 0))


def _merged(staff: list[dict], roles: set[str], by_role: bool = False) -> list[tuple]:
    """One row per person (per person and role if by_role), summing hours, in input order."""
    hours: dict[tuple, Decimal] = {}
    for s in staff:
        if s["role"] in roles and _d(s["hours"]) > 0:
            key = (s["name"], s["role"]) if by_role else (s["name"],)
            hours[key] = hours.get(key, Decimal("0")) + _d(s["hours"])
    return [(*key, h) for key, h in hours.items()]


def form_math(p: dict, support: list, tipped: list) -> dict:
    """What the form's formulas will calculate, mirroring CashOut_Form.xlsx cell by cell."""
    pool = _d(p.get("auto_gratuity")) + _d(p.get("pool_card_tips")) + _d(p.get("cash_tips_manual"))  # B16
    kitchen = Decimal("0.1") * _d(p.get("gross_food_sales"))  # B19
    roles = {role for _, role, h in support if h > 0}  # B20: how many different support roles have hours
    rate = Decimal("0.0225") if len(roles) >= 2 else Decimal("0.015") if roles else Decimal("0")
    support_total = rate * _d(p.get("net_sales"))  # B21
    remainder = pool - kitchen - support_total  # B26
    s_hours = sum((h for _, _, h in support), Decimal("0"))
    t_hours = sum((h for _, h in tipped), Decimal("0"))
    payouts: dict[str, Decimal] = {}
    for name, _, h in support:  # D23:D25
        payouts[name] = payouts.get(name, Decimal("0")) + (support_total * h / s_hours if s_hours else Decimal("0"))
    for name, h in tipped:  # E30:E39
        payouts[name] = payouts.get(name, Decimal("0")) + (remainder * h / t_hours if t_hours else Decimal("0"))
    return {"kitchen": kitchen, "support_total": support_total, "payouts": payouts}


def check_period(name: str, p: dict, result) -> tuple[list, list, list[str]]:
    support = _merged(p.get("staff", []), SUPPORT_ROLES, by_role=True)
    tipped = _merged(p.get("staff", []), TIPPED_ROLES)
    problems = []
    if len(support) > len(SUPPORT_ROWS):
        problems.append(f"{name}: {len(support)} support staff worked but the form only has {len(SUPPORT_ROWS)} rows.")
    if len(tipped) > len(TIPPED_ROWS):
        problems.append(f"{name}: {len(tipped)} bartenders/servers worked but the form only has {len(TIPPED_ROWS)} rows.")
    if problems:
        return support, tipped, problems

    form = form_math(p, support, tipped)
    engine = {row[0]: Decimal("0") for row in support + tipped}
    for n, amt in result.support_payouts.items():
        engine[n] += amt
    for n, amt in result.tipped_payouts.items():
        engine[n] += amt
    if abs(form["kitchen"] - result.kitchen) >= TOLERANCE:
        problems.append(f"{name}: form kitchen {form['kitchen']:.2f} vs engine {result.kitchen}.")
    if abs(form["support_total"] - result.support_total) >= TOLERANCE:
        problems.append(
            f"{name}: form support {form['support_total']:.2f} vs engine {result.support_total}."
        )
    for person in engine:
        if abs(form["payouts"][person] - engine[person]) >= TOLERANCE:
            problems.append(f"{name}: {person} would get {form['payouts'][person]:.2f} on the form vs {engine[person]} from the engine.")
    if result.due_back > 0:
        problems.append(f"{name}: DUE BACK ${result.due_back} - the form would show negative payouts for bar/servers.")
    return support, tipped, problems


def fill(doc: dict, out: Path, completed_by: str | None) -> dict:
    day = compute_day(doc)
    periods = doc.get("periods", {})
    results = {r.period: r for r in day["periods"]}

    checked = {}
    problems: list[str] = []
    for name, p in periods.items():
        support, tipped, period_problems = check_period(name, p, results[name])
        checked[name] = (support, tipped)
        problems += period_problems
    if problems:
        return {"written": None, "problems": problems}

    wb = load_workbook(FORM)
    template = wb[TEMPLATE_SHEET]
    for name, p in periods.items():
        support, tipped = checked[name]
        ws = wb.copy_worksheet(template)
        ws.title = name
        ws["B4"] = doc.get("business_date")
        ws["E4"] = name
        if completed_by:
            ws["B5"] = completed_by
        if p.get("cash_sales") is not None:
            ws["B8"] = float(p["cash_sales"])
        ws["B9"] = float(_d(p.get("net_sales")))
        ws["B10"] = float(_d(p.get("gross_food_sales")))
        ws["B13"] = float(_d(p.get("auto_gratuity")))
        ws["B14"] = float(_d(p.get("pool_card_tips")))
        ws["B15"] = float(_d(p.get("cash_tips_manual")))
        for row, (person, role, hours) in zip(SUPPORT_ROWS, support):
            ws[f"B{row}"], ws[f"C{row}"], ws[f"E{row}"] = person, float(hours), role
        for row, (person, hours) in zip(TIPPED_ROWS, tipped):
            ws[f"B{row}"], ws[f"C{row}"] = person, float(hours)
        # copy_worksheet doesn't carry data validation; restore the dropdowns.
        for cells, choices in (("E4", '"Lunch,Dinner"'), ("E23:E25", '"Host,Barback,Busser,Runner"')):
            dv = DataValidation(type="list", formula1=choices, allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(cells)
        ws.print_options.horizontalCentered = template.print_options.horizontalCentered
        ws.page_setup.orientation = template.page_setup.orientation
        ws.page_setup.fitToWidth = template.page_setup.fitToWidth
        ws.sheet_properties.pageSetUpPr = template.sheet_properties.pageSetUpPr
    for sheet in [s for s in wb.sheetnames if s not in periods]:
        del wb[sheet]
    wb.save(out)
    return {
        "written": str(out),
        "problems": [],
        "sheets": list(periods),
        "per_person": {n: str(a) for n, a in day["per_person"].items()},
        "kitchen_lump_total": str(day["kitchen_lump_total"]),
        "hourly_rate": {r.period: str(r.hourly_rate) for r in day["periods"]},
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("input")
    parser.add_argument("--out", required=True)
    parser.add_argument("--completed-by")
    args = parser.parse_args()
    summary = fill(json.loads(Path(args.input).read_text()), Path(args.out), args.completed_by)
    print(json.dumps(summary, indent=2))
    return 0 if summary["written"] else 2


if __name__ == "__main__":
    sys.exit(main())
